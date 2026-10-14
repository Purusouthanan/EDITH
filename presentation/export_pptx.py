# -*- coding: utf-8 -*-
"""
EDITH Hackathon Pitch Deck — Complete PowerPoint (PPTX) Generator
Creates: EDITH_Hackathon_Pitch_Deck.pptx (16:9 Widescreen, Fully Editable)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

WORKSPACE = r"c:\Users\PURUSOUTHANAN\OneDrive\Desktop\EDITH"
PPTX_PATH = os.path.join(WORKSPACE, "EDITH_Hackathon_Pitch_Deck.pptx")

# Color Palette
BG_COLOR = RGBColor(9, 16, 32)        # #091020 Deep Navy
CARD_BG = RGBColor(17, 29, 51)        # #111D33 Card Navy
CARD_BORDER = RGBColor(76, 145, 255)  # #4C91FF Blue Accent
CYAN = RGBColor(43, 220, 224)         # #2BDCE0 Cyan Highlight
BLUE = RGBColor(76, 145, 255)         # #4C91FF Blue
GREEN = RGBColor(82, 220, 164)        # #52DCA4 Green Verified
ORANGE = RGBColor(255, 164, 76)       # #FFA44C Orange Warning
RED = RGBColor(255, 104, 120)         # #FF6878 Red Critical
WHITE = RGBColor(244, 248, 255)       # #F4F8FF Crisp White
TEXT_MUTED = RGBColor(163, 180, 204)  # #A3B4CC Light Slate

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def set_slide_background(slide):
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_COLOR
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, eyebrow_text, title_text, subtitle_text, slide_num):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p1 = tf.paragraphs[0]
    p1.text = f"● {eyebrow_text.upper()}"
    p1.font.size = Pt(9.5)
    p1.font.bold = True
    p1.font.color.rgb = CYAN
    p1.font.name = "Arial"
    
    p2 = tf.add_paragraph()
    p2.text = title_text
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.font.name = "Arial"
    p2.space_before = Pt(3)
    
    p3 = tf.add_paragraph()
    p3.text = subtitle_text
    p3.font.size = Pt(11)
    p3.font.color.rgb = TEXT_MUTED
    p3.font.name = "Arial"
    p3.space_before = Pt(2)
    
    badge_box = slide.shapes.add_textbox(Inches(10.8), Inches(0.4), Inches(1.7), Inches(0.35))
    btf = badge_box.text_frame
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.RIGHT
    bp.text = f"SLIDE {slide_num:02d} / 16"
    bp.font.size = Pt(9.5)
    bp.font.bold = True
    bp.font.color.rgb = BLUE
    bp.font.name = "Arial"

def add_footer(slide, section_name):
    fb = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(11.733), Inches(0.35))
    ftf = fb.text_frame
    ftf.word_wrap = True
    ftf.margin_left = tf_margin = 0
    p = ftf.paragraphs[0]
    p.text = f"EDITH · EMERGENCY DISPATCH & INTELLIGENCE TRACKING HUB   |   TEAM JARVIS · RATHINAM TECHNICAL CAMPUS   |   {section_name}"
    p.font.size = Pt(8.5)
    p.font.color.rgb = RGBColor(114, 132, 158)
    p.font.name = "Arial"

def add_card(slide, left, top, width, height, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1)
    return shape

# ==========================================
# SLIDE 1: COVER
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1)

tb = s1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.5))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "● HEALTHCARE DISASTER COORDINATION PLATFORM · FLOOD & CYCLONE EMERGENCY RESPONSE"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = CYAN
p.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = "EDITH"
p2.font.size = Pt(64)
p2.font.bold = True
p2.font.color.rgb = WHITE
p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(16)

p3 = tf.add_paragraph()
p3.text = "Emergency Dispatch & Intelligence Tracking Hub"
p3.font.size = Pt(22)
p3.font.bold = True
p3.font.color.rgb = BLUE
p3.alignment = PP_ALIGN.CENTER
p3.space_before = Pt(4)

p4 = tf.add_paragraph()
p4.text = '"From fragmented emergency updates to accountable healthcare coordination."'
p4.font.size = Pt(15)
p4.font.italic = True
p4.font.color.rgb = CYAN
p4.alignment = PP_ALIGN.CENTER
p4.space_before = Pt(12)

# 4 Cards on Cover
card_w = Inches(2.7)
card_h = Inches(1.5)
top_y = Inches(4.3)

c_data = [
    ("Offline-First Outbox", "Service Worker + IndexedDB queue preserves field requests during blackouts.", CYAN),
    ("Request-Centric Model", "Facilities volunteer surplus on demand; zero mandatory inventory exposure.", BLUE),
    ("Atomic Allocation", "PostgreSQL ACID transactions eliminate over-allocation in race conditions.", GREEN),
    ("Two-Phase Verification", "Driver drop-off does not equal delivery; recipient doctor confirms receipt.", ORANGE)
]

for idx, (title, desc, col) in enumerate(c_data):
    cx = Inches(1.0 + idx * 2.9)
    add_card(s1, cx, top_y, card_w, card_h, col)
    tb_c = s1.shapes.add_textbox(cx + Inches(0.15), top_y + Inches(0.15), card_w - Inches(0.3), card_h - Inches(0.3))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p_t = tf_c.paragraphs[0]
    p_t.text = title
    p_t.font.bold = True
    p_t.font.size = Pt(12)
    p_t.font.color.rgb = col
    p_d = tf_c.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = TEXT_MUTED
    p_d.space_before = Pt(4)

# Footer on Cover
add_footer(s1, "TEAM JARVIS: Purusouthanan B (Lead) · Sakthivel R · Sanjo Thomas · Sheik Arsh S")

# ==========================================
# SLIDE 2: MEET TEAM JARVIS
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2)
add_header(s2, "SECTION A · TEAM INTRODUCTION", "Meet Team JARVIS · Rathinam Technical Campus", "Multidisciplinary student engineering team dedicated to solving disaster-zone coordination breakdown.", 2)

members = [
    ("Purusouthanan B", "Team Lead · Systems Architect", "Leads end-to-end architectural design, disaster workflow state machines, role-based boundary modeling, and technical planning.", CYAN),
    ("Sakthivel R", "Backend & API Engineer", "Engineers FastAPI modular monolith services, Pydantic schema validation contracts, RESTful endpoint lifecycle, and error handlers.", BLUE),
    ("Sanjo Thomas", "Frontend & Offline Systems", "Builds lightweight PWA interfaces, Service Worker cache strategies, IndexedDB transaction queues, and offline synchronization managers.", GREEN),
    ("Sheik Arsh S", "Data Integrity & QA Specialist", "Oversees PostgreSQL relational constraints, row-level concurrency locks, idempotency verification, and chaos testing scenarios.", ORANGE)
]

m_w = Inches(2.75)
m_h = Inches(3.4)
m_top = Inches(1.8)

for idx, (name, role, bio, col) in enumerate(members):
    mx = Inches(0.8 + idx * 3.0)
    add_card(s2, mx, m_top, m_w, m_h, col)
    tb_m = s2.shapes.add_textbox(mx + Inches(0.18), m_top + Inches(0.18), m_w - Inches(0.36), m_h - Inches(0.36))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    
    p_n = tf_m.paragraphs[0]
    p_n.text = name
    p_n.font.bold = True
    p_n.font.size = Pt(15)
    p_n.font.color.rgb = WHITE
    
    p_r = tf_m.add_paragraph()
    p_r.text = role
    p_r.font.bold = True
    p_r.font.size = Pt(10.5)
    p_r.font.color.rgb = col
    p_r.space_before = Pt(3)
    
    p_b = tf_m.add_paragraph()
    p_b.text = bio
    p_b.font.size = Pt(10)
    p_b.font.color.rgb = TEXT_MUTED
    p_b.space_before = Pt(8)

# Bottom philosophy card
p_box = add_card(s2, Inches(0.8), Inches(5.4), Inches(11.733), Inches(1.3), CYAN)
tb_p = s2.shapes.add_textbox(Inches(1.0), Inches(5.45), Inches(11.333), Inches(1.2))
tf_p = tb_p.text_frame
tf_p.word_wrap = True
p_pt = tf_p.paragraphs[0]
p_pt.text = "TEAM JARVIS FOUNDATION PILLARS"
p_pt.font.bold = True
p_pt.font.size = Pt(10)
p_pt.font.color.rgb = CYAN
p_pd = tf_p.add_paragraph()
p_pd.text = "1. Ground-Truth Analysis: Studied failure modes in Tamil Nadu cyclone/flood relief.  2. Deliberate Simplicity: Avoided fragile microservices in favor of a clean modular monolith.  3. Zero False Guarantees: Never misrepresent un-synced offline drafts as confirmed server records.  4. Non-Repudiation: Every state shift has an immutable human actor record."
p_pd.font.size = Pt(9.5)
p_pd.font.color.rgb = WHITE
p_pd.space_before = Pt(4)

add_footer(s2, "SECTION A: TEAM INTRODUCTION")

# ==========================================
# SLIDE 3: WHY THIS PROBLEM MATTERS
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3)
add_header(s3, "SECTION A · OPERATIONAL CONTEXT", "Why This Problem Matters: Healthcare in Catastrophic Disasters", "In severe floods and cyclones, the bottleneck is rarely total supply absence, but a failure of coordination.", 3)

c_w = Inches(3.7)
c_h = Inches(3.3)
top_y = Inches(1.8)

issues = [
    ("1. Demand Surges & Depletion", "Floods and cyclones cause sudden spikes in trauma, waterborne illness, and power cuts. Oxygen cylinders, O- blood, tetanus toxoids, and generator diesel exhaust within hours at isolated medical camps.", RED),
    ("2. Fragmented Channel Chaos", "Requests broadcast across disconnected WhatsApp groups, phone calls, and paper memos. A high-profile camp gets 4 duplicate donations while an isolated camp 5 km away starves.", ORANGE),
    ("3. Critical Infrastructure Blackout", "Downed cell towers, flooded electrical substations, and cut fiber backhauls disconnect cloud services for days. Standard cloud apps and EHRs freeze instantly upon signal loss.", BLUE)
]

for idx, (title, desc, col) in enumerate(issues):
    cx = Inches(0.8 + idx * 4.0)
    add_card(s3, cx, top_y, c_w, c_h, col)
    tb_c = s3.shapes.add_textbox(cx + Inches(0.2), top_y + Inches(0.2), c_w - Inches(0.4), c_h - Inches(0.4))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    pt = tf_c.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(14)
    pt.font.color.rgb = WHITE
    pd = tf_c.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(11)
    pd.font.color.rgb = TEXT_MUTED
    pd.space_before = Pt(8)

# Horizontal breakdown banner
b_box = add_card(s3, Inches(0.8), Inches(5.35), Inches(11.733), Inches(1.35), CYAN)
tb_b = s3.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(11.333), Inches(1.2))
tf_b = tb_b.text_frame
tf_b.word_wrap = True
p_bt = tf_b.paragraphs[0]
p_bt.text = "ANATOMY OF DISASTER COORDINATION COLLAPSE"
p_bt.font.bold = True
p_bt.font.size = Pt(10)
p_bt.font.color.rgb = CYAN
p_bd = tf_b.add_paragraph()
p_bd.text = "Surge Demand Spike  ──>  Multi-Channel WhatsApp Shouting  ──>  Unacknowledged Requests  ──>  Blind Unverified Dispatches  ──>  Zero Proof of Delivery\nCore Takeaway: Software designed exclusively for always-on 5G links fails immediately when power and telecom towers fail."
p_bd.font.size = Pt(10)
p_bd.font.color.rgb = WHITE
p_bd.space_before = Pt(4)

add_footer(s3, "SECTION A: OPERATIONAL CONTEXT")

# ==========================================
# SLIDE 4: THE PROBLEM STATEMENT
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4)
add_header(s4, "SECTION B · PROBLEM STATEMENT", "The Problem Statement: Coordination Without Accountability", "Formalizing the critical gap in emergency healthcare response during severe weather disasters.", 4)

# Big Formal Problem Box
p_box = add_card(s4, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.6), CYAN)
tb_pb = s4.shapes.add_textbox(Inches(1.1), Inches(1.95), Inches(11.133), Inches(1.3))
tf_pb = tb_pb.text_frame
tf_pb.word_wrap = True
p_pbt = tf_pb.paragraphs[0]
p_pbt.text = "FORMAL HACKATHON PROBLEM STATEMENT"
p_pbt.font.bold = True
p_pbt.font.size = Pt(10)
p_pbt.font.color.rgb = CYAN
p_pbd = tf_pb.add_paragraph()
p_pbd.text = '"During severe floods and cyclones, hospitals and medical camps struggle to coordinate emergency resource requests because communication channels are fragmented, availability data becomes stale, ownership is unclear, and network disruptions interrupt cloud-dependent workflows."'
p_pbd.font.size = Pt(13)
p_pbd.font.bold = True
p_pbd.font.italic = True
p_pbd.font.color.rgb = WHITE
p_pbd.space_before = Pt(4)

# 5 Pain Points
pain_w = Inches(2.2)
pain_h = Inches(3.0)
pain_top = Inches(3.65)

pains = [
    ("Fragmented Comm", "Phone calls, SMS, WhatsApp, and paper create disconnected silos with no unified record.", CYAN),
    ("Stale Data", "Spreadsheets and verbal promises don't reflect rapid field depletion or prior commitments.", BLUE),
    ("Unclear Ownership", "Urgent needs entered into group chats sit unacknowledged without an assigned coordinator.", ORANGE),
    ("Unverified Fulfilment", "'Dispatched' is conflated with 'Delivered'. Lost, delayed, or stolen cargo goes undetected.", RED),
    ("Cloud Fragility", "Standard web portals freeze or crash the moment local cell towers lose grid power.", GREEN)
]

for idx, (title, desc, col) in enumerate(pains):
    px = Inches(0.8 + idx * 2.4)
    add_card(s4, px, pain_top, pain_w, pain_h, col)
    tb_pain = s4.shapes.add_textbox(px + Inches(0.15), pain_top + Inches(0.15), pain_w - Inches(0.3), pain_h - Inches(0.3))
    tf_pain = tb_pain.text_frame
    tf_pain.word_wrap = True
    pt = tf_pain.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(12)
    pt.font.color.rgb = col
    pd = tf_pain.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = TEXT_MUTED
    pd.space_before = Pt(6)

add_footer(s4, "SECTION B: PROBLEM STATEMENT")

# ==========================================
# SLIDE 5: EXISTING GAP VS EDITH (TABLE)
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5)
add_header(s5, "SECTION B · COMPARATIVE ANALYSIS", "Existing Coordination Gap vs. The EDITH Standard", "Moving from informal, fragile communication to structured, resilient accountability.", 5)

rows = 9
cols = 4
tbl_left = Inches(0.8)
tbl_top = Inches(1.8)
tbl_w = Inches(11.733)
tbl_h = Inches(4.8)

tbl_shape = s5.shapes.add_table(rows, cols, tbl_left, tbl_top, tbl_w, tbl_h)
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.2)
tbl.columns[1].width = Inches(2.8)
tbl.columns[2].width = Inches(2.8)
tbl.columns[3].width = Inches(3.933)

headers = ["Dimension", "Informal (Phone/WhatsApp)", "Generic Dashboards / Excel", "EDITH Platform Standard"]
for c_idx, h in enumerate(headers):
    cell = tbl.cell(0, c_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = CARD_BG
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = CYAN if c_idx == 3 else WHITE

table_data = [
    ("Request Visibility", "Siloed in private chats", "Passive tabular broadcast", "Unified regional triage queue with priority tags"),
    ("Ownership & SLA", "No assigned coordinator", "Ambiguous ticket states", "Mandatory ICC coordinator assignment with timers"),
    ("Resource Offers", "Verbal chat promises", "Static passive inventory", "Structured voluntary offers tied to active needs"),
    ("Allocation Tracking", "High double-promising risk", "Manual cell edits (race risk)", "ACID row-locks preventing over-allocation"),
    ("Partial Fulfilment", "Lost in chat history", "Binary complete/incomplete", "Automated state split: fulfilled vs remaining"),
    ("Delivery Confirmation", "Assumed on driver call", "Driver click without proof", "Two-phase: Recipient doctor verifies count"),
    ("Offline Operation", "Fails when signal drops", "Browser errors & lost forms", "Service Worker + IndexedDB outbox queue"),
    ("Auditability & History", "Ephemeral, easily deleted", "Basic edit timestamps", "Immutable append-only audit trail with diffs")
]

for r_idx, r_data in enumerate(table_data):
    for c_idx, val in enumerate(r_data):
        cell = tbl.cell(r_idx + 1, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(14, 24, 43) if r_idx % 2 == 0 else CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.size = Pt(9.5)
        if c_idx == 3:
            p.font.color.rgb = GREEN
            p.font.bold = True
        elif c_idx == 0:
            p.font.color.rgb = WHITE
            p.font.bold = True
        else:
            p.font.color.rgb = TEXT_MUTED

add_footer(s5, "SECTION B: COMPARATIVE ANALYSIS")

# ==========================================
# SLIDE 6: PROPOSED SOLUTION WORKFLOW
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6)
add_header(s6, "SECTION B · PROPOSED SOLUTION", "The EDITH Operational Lifecycle: 7-Stage Accountable Coordination", "Every emergency request moves through an explicit, state-machine-governed progression with non-repudiation.", 6)

steps = [
    ("1. SUBMITTED", "Camp logs urgent need (e.g. 20 O2 cylinders)", CYAN),
    ("2. ACKNOWLEDGED", "ICC coordinator takes operational ownership", BLUE),
    ("3. OFFER MADE", "Nearby hospital volunteers surplus (12 units)", CYAN),
    ("4. ASSIGNED", "ICC approves allocation; DB row-lock committed", ORANGE),
    ("5. IN TRANSIT", "Logistics packages cargo, logs vehicle ID & dispatches", BLUE),
    ("6. DROP-OFF", "Driver logs arrival at camp gate (Drop-off only)", ORANGE),
    ("7. CONFIRMED", "Recipient physician inspects & verifies quantity", GREEN)
]

s_w = Inches(1.55)
s_h = Inches(2.6)
s_top = Inches(1.8)

for idx, (title, desc, col) in enumerate(steps):
    sx = Inches(0.8 + idx * 1.7)
    add_card(s6, sx, s_top, s_w, s_h, col)
    tb_s = s6.shapes.add_textbox(sx + Inches(0.1), s_top + Inches(0.12), s_w - Inches(0.2), s_h - Inches(0.24))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    pt = tf_s.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(10)
    pt.font.color.rgb = col
    pd = tf_s.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(8.5)
    pd.font.color.rgb = TEXT_MUTED
    pd.space_before = Pt(6)

# 3 Exceptions Cards
ex_w = Inches(3.7)
ex_h = Inches(1.9)
ex_top = Inches(4.7)

exceptions = [
    ("Exception A: Partial Fulfilment Split", "If 12 of 20 units are verified, 12 units resolve while 8 units remain active in the regional triage queue for a secondary donor.", ORANGE),
    ("Exception B: Delivery Dispute Protocol", "If cargo arrives damaged or missing, recipient flags dispute. Resolution halts, ICC alerts trigger, and an audit review begins.", RED),
    ("Exception C: Reassignment & Road Pivots", "If floodwaters submerge transit roads, coordinators can reassign cargo to alternate staging hubs with complete historical logging.", CYAN)
]

for idx, (title, desc, col) in enumerate(exceptions):
    ex = Inches(0.8 + idx * 4.0)
    add_card(s6, ex, ex_top, ex_w, ex_h, col)
    tb_e = s6.shapes.add_textbox(ex + Inches(0.15), ex_top + Inches(0.15), ex_w - Inches(0.3), ex_h - Inches(0.3))
    tf_e = tb_e.text_frame
    tf_e.word_wrap = True
    pt = tf_e.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(11)
    pt.font.color.rgb = col
    pd = tf_e.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = TEXT_MUTED
    pd.space_before = Pt(4)

add_footer(s6, "SECTION B: PROPOSED SOLUTION")

# ==========================================
# SLIDE 7: OFFLINE RESILIENCE
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7)
add_header(s7, "SECTION B · OFFLINE RESILIENCE", "Offline Resilience: The EDITH Difference", "Preserving operational continuity during network blackouts — without corrupting central state.", 7)

seq = [
    ("Step 1: Online Baseline", "Standard HTTPS queries to FastAPI backend.", GREEN),
    ("Step 2: Signal Drops", "Service worker catches dropped calls gracefully.", RED),
    ("Step 3: Local Capture", "Request saved immediately to IndexedDB outbox.", ORANGE),
    ("Step 4: Explicit UI Truth", "Displays 'Queued Locally · Pending Sync' amber badge.", CYAN),
    ("Step 5: Signal Restored", "Sync manager replays payload with Idempotency Key.", BLUE),
    ("Step 6: Server Commit", "FastAPI validates & commits to PostgreSQL truth store.", GREEN)
]

sq_w = Inches(1.8)
sq_h = Inches(2.2)
sq_top = Inches(1.8)

for idx, (title, desc, col) in enumerate(seq):
    sqx = Inches(0.8 + idx * 2.0)
    add_card(s7, sqx, sq_top, sq_w, sq_h, col)
    tb_sq = s7.shapes.add_textbox(sqx + Inches(0.1), sq_top + Inches(0.12), sq_w - Inches(0.2), sq_h - Inches(0.24))
    tf_sq = tb_sq.text_frame
    tf_sq.word_wrap = True
    pt = tf_sq.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(10.5)
    pt.font.color.rgb = col
    pd = tf_sq.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(9)
    pd.font.color.rgb = TEXT_MUTED
    pd.space_before = Pt(4)

# 3 Engineering Principles
p_w = Inches(3.7)
p_h = Inches(2.3)
p_top = Inches(4.3)

principles = [
    ("No False Mesh Claims", "Separate devices cannot magically sync across floodwaters without an active physical or radio link. EDITH never makes false claims of peer-to-peer miracles.", CYAN),
    ("Truth Transparency", "A locally-saved draft is never presented to the user as a server-confirmed record. Un-synced requests carry a distinct visual warning and temporary UUID.", ORANGE),
    ("Idempotency Protection", "Repeated sync retries on flickering 2G links cannot create duplicate tickets or double-allocate resources. Server idempotency keys enforce safe replay.", GREEN)
]

for idx, (title, desc, col) in enumerate(principles):
    px = Inches(0.8 + idx * 4.0)
    add_card(s7, px, p_top, p_w, p_h, col)
    tb_pr = s7.shapes.add_textbox(px + Inches(0.18), p_top + Inches(0.18), p_w - Inches(0.36), p_h - Inches(0.36))
    tf_pr = tb_pr.text_frame
    tf_pr.word_wrap = True
    pt = tf_pr.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(12)
    pt.font.color.rgb = col
    pd = tf_pr.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(10)
    pd.font.color.rgb = TEXT_MUTED
    pd.space_before = Pt(6)

add_footer(s7, "SECTION B: OFFLINE RESILIENCE")

# ==========================================
# SLIDE 8: SYSTEM ARCHITECTURE
# ==========================================
s8 = prs.slides.add_slide(blank_layout)
set_slide_background(s8)
add_header(s8, "SECTION C · SYSTEM ARCHITECTURE", "Three-Tier Modular Architecture for Resilient Coordination", "Strict separation of concerns enforcing human authority, offline survival, and ACID truth.", 8)

tiers = [
    ("TIER 1 · CLIENT PWA", "Vanilla HTML5 + CSS3 + Modern ES6+ JS\n\n• Service Worker: Application shell caching & offline fallback\n• IndexedDB: Outbox queue for mutations & read cache\n• Sync Engine: Background retry with exponential backoff\n• Target: <500ms initial load on 2G connections", CYAN),
    ("TIER 2 · API SERVICES", "FastAPI (v0.115+) Modular Monolith\n\n• Modules: Auth (JWT/RBAC), Requests, Offers, Allocations, Sync\n• Pydantic v2: Strict contract validation on offline payloads\n• Idempotency Guard: Deduplication middleware via UUID keys\n• Performance: Asynchronous non-blocking event loop", BLUE),
    ("TIER 3 · DATA STORE", "PostgreSQL 16 (Dockerized)\n\n• ACID Guarantees: Authoritative truth with relational schema\n• Row-Level Locks: SELECT ... FOR UPDATE on allocations\n• Append-Only Audit Log: Immutable table logging all diffs\n• Deployment: Single-container reproducible image", GREEN)
]

t_w = Inches(3.7)
t_h = Inches(3.6)
t_top = Inches(1.8)

for idx, (title, desc, col) in enumerate(tiers):
    tx = Inches(0.8 + idx * 4.0)
    add_card(s8, tx, t_top, t_w, t_h, col)
    tb_t = s8.shapes.add_textbox(tx + Inches(0.2), t_top + Inches(0.2), t_w - Inches(0.4), t_h - Inches(0.4))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    pt = tf_t.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(14)
    pt.font.color.rgb = col
    pd = tf_t.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(10.5)
    pd.font.color.rgb = WHITE
    pd.space_before = Pt(8)

# AI Integration Banner
ai_box = add_card(s8, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.1), ORANGE)
tb_ai = s8.shapes.add_textbox(Inches(1.0), Inches(5.65), Inches(11.333), Inches(1.0))
tf_ai = tb_ai.text_frame
tf_ai.word_wrap = True
pt = tf_ai.paragraphs[0]
pt.text = "AUXILIARY INTEGRATION: OPTIONAL ASSISTIVE AI SCOPE"
pt.font.bold = True
pt.font.size = Pt(10)
pt.font.color.rgb = ORANGE
pd = tf_ai.add_paragraph()
pd.text = "External LLM APIs strictly limited to assistive tasks: audio transcription, structured request extraction, and triage summaries. Zero autonomous decision power; fail-safe graceful degradation if offline."
pd.font.size = Pt(9.5)
pd.font.color.rgb = WHITE
pd.space_before = Pt(3)

add_footer(s8, "SECTION C: SYSTEM ARCHITECTURE")

# ==========================================
# SLIDE 9: TECHNOLOGY MATRIX (TABLE)
# ==========================================
s9 = prs.slides.add_slide(blank_layout)
set_slide_background(s9)
add_header(s9, "SECTION C · TECHNOLOGY MATRIX", "Deliberate Engineering Choices: Built for Field Feasibility", "Every technology selected for minimal resource footprint, high resilience, and rapid deployment.", 9)

tbl_shape9 = s9.shapes.add_table(9, 3, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
tbl9 = tbl_shape9.table
tbl9.columns[0].width = Inches(2.2)
tbl9.columns[1].width = Inches(3.2)
tbl9.columns[2].width = Inches(6.333)

h9 = ["Architectural Layer", "Selected Technology", "Operational Justification & Field Fit"]
for c_idx, h in enumerate(h9):
    cell = tbl9.cell(0, c_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = CARD_BG
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = CYAN

stack_data = [
    ("Frontend Interface", "Vanilla HTML5, CSS3, Modern JS", "Zero framework compilation overhead; instant load even on budget smartphones."),
    ("Offline Shell", "Service Worker & Cache API", "Enables offline launching, home-screen installation without app store friction."),
    ("Client Persistence", "IndexedDB (Native Browser DB)", "Structured local database for outbox queue and read cache; survives browser reboots."),
    ("Backend Framework", "Python FastAPI (v0.115+)", "Lightweight ASGI micro-framework with native async pipeline & auto OpenAPI schema."),
    ("Contract Validation", "Pydantic v2 (Rust Core)", "High-speed data boundary validation catching corrupted offline payloads cleanly."),
    ("Data Access / ORM", "SQLAlchemy 2.0 (Psycopg 3)", "Type-safe ORM with explicit transaction control and row-level locking support."),
    ("Authoritative Store", "PostgreSQL 16 (Dockerized)", "World-class ACID durability, proven concurrency control, and JSONB audit trails."),
    ("Testing & QA", "Pytest, Asyncio, Playwright", "Automated test suites covering concurrent allocations and offline sync replays.")
]

for r_idx, (layer, tech, just) in enumerate(stack_data):
    row_bg = RGBColor(14, 24, 43) if r_idx % 2 == 0 else CARD_BG
    for c_idx, val in enumerate([layer, tech, just]):
        cell = tbl9.cell(r_idx + 1, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = row_bg
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.size = Pt(9.5)
        if c_idx == 0:
            p.font.bold = True
            p.font.color.rgb = WHITE
        elif c_idx == 1:
            p.font.bold = True
            p.font.color.rgb = CYAN if "HTML" in val or "Service" in val or "IndexedDB" in val else BLUE
        else:
            p.font.color.rgb = TEXT_MUTED

add_footer(s9, "SECTION C: TECHNOLOGY MATRIX")

# ==========================================
# SLIDE 10: INTEGRITY & SECURITY
# ==========================================
s10 = prs.slides.add_slide(blank_layout)
set_slide_background(s10)
add_header(s10, "SECTION C · INTEGRITY & SECURITY", "Data Integrity, Security & Concurrency Controls", "Engineering safeguards that eliminate race conditions, unauthorized actions, and data corruption.", 10)

sec_w = Inches(5.7)
sec_h = Inches(2.2)

sec_items = [
    ("1. Role-Based Access Control (RBAC)", "• Hospital / Camp Users: Scoped strictly to drafting requests & offers for their facility.\n• ICC Coordinators: Global zone authority to acknowledge, triage, allocate, and close.\n• Logistics Drivers: Scoped strictly to updating transit status and logging drop-offs.", CYAN, Inches(0.8), Inches(1.8)),
    ("2. Atomic Allocation & Row Locking", "• Problem: Two coordinators simultaneously assign the last 10 oxygen cylinders.\n• Solution: SELECT quantity_available FROM offers WHERE id = :id FOR UPDATE\n• Result: Transaction A locks & commits; Transaction B detects 0 and rejects safely.", GREEN, Inches(6.8), Inches(1.8)),
    ("3. Idempotency & Duplicate Prevention", "• Offline mutations attach a client-generated UUID v4 idempotency key.\n• Server checks processed_operations table before execution.\n• Network retries on flickering links return cached results without double records.", ORANGE, Inches(0.8), Inches(4.3)),
    ("4. Two-Phase Fulfilment Verification", "• Driver drop-off logs physical arrival with timestamp at the camp gate.\n• Resolution is NOT triggered: Status becomes DELIVERED_PENDING_VERIFICATION.\n• The receiving doctor must log in, inspect cargo, and verify physical count.", RED, Inches(6.8), Inches(4.3))
]

for title, desc, col, lx, ty in sec_items:
    add_card(s10, lx, ty, sec_w, sec_h, col)
    tb_s = s10.shapes.add_textbox(lx + Inches(0.2), ty + Inches(0.18), sec_w - Inches(0.4), sec_h - Inches(0.36))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    pt = tf_s.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(12)
    pt.font.color.rgb = col
    pd = tf_s.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = WHITE
    pd.space_before = Pt(4)

add_footer(s10, "SECTION C: DATA INTEGRITY & SECURITY")

# ==========================================
# SLIDE 11: COMMAND DASHBOARD (DEMO)
# ==========================================
s11 = prs.slides.add_slide(blank_layout)
set_slide_background(s11)
add_header(s11, "SECTION D · INTERFACE MOCKUP (ILLUSTRATIVE DEMO DATA)", "Command-Centre Dashboard: Real-Time Operational Triage", "High-density, low-cognitive-load interface designed for disaster coordinators.", 11)

# Top ribbon
r_box = add_card(s11, Inches(0.8), Inches(1.7), Inches(11.733), Inches(0.55), BLUE)
tb_r = s11.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(11.333), Inches(0.45))
tf_r = tb_r.text_frame
p_r = tf_r.paragraphs[0]
p_r.text = "INCIDENT COMMAND CENTRE · SECTOR 4 DISASTER DESK   |   ● NETWORK: CONNECTED   |   LAST SYNC: 3s AGO   |   SIMULATED TEST DATA"
p_r.font.size = Pt(9.5)
p_r.font.bold = True
p_r.font.color.rgb = WHITE

# 4 KPI cards
k_w = Inches(2.75)
k_h = Inches(1.1)
k_top = Inches(2.4)
kpis = [
    ("TOTAL OPEN REQUESTS", "34", "8 Critical · 14 High · 12 Normal", CYAN),
    ("UNACKNOWLEDGED NEEDS", "3", "Immediate Triage Required", RED),
    ("OFFERS PENDING SIGN-OFF", "6", "Matches Awaiting Allocation", ORANGE),
    ("IN TRANSIT / ACTIVE FLEET", "12", "Logistics En Route", GREEN)
]

for idx, (lbl, val, sub, col) in enumerate(kpis):
    kx = Inches(0.8 + idx * 3.0)
    add_card(s11, kx, k_top, k_w, k_h, col)
    tb_k = s11.shapes.add_textbox(kx + Inches(0.15), k_top + Inches(0.1), k_w - Inches(0.3), k_h - Inches(0.2))
    tf_k = tb_k.text_frame
    tf_k.word_wrap = True
    p1 = tf_k.paragraphs[0]
    p1.text = lbl
    p1.font.size = Pt(8.5)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_MUTED
    p2 = tf_k.add_paragraph()
    p2.text = val
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = col
    p3 = tf_k.add_paragraph()
    p3.text = sub
    p3.font.size = Pt(8.5)
    p3.font.color.rgb = WHITE

# Left: Table
tbl_shape11 = s11.shapes.add_table(5, 5, Inches(0.8), Inches(3.7), Inches(7.5), Inches(2.9))
tbl11 = tbl_shape11.table
tbl11.columns[0].width = Inches(1.1)
tbl11.columns[1].width = Inches(2.0)
tbl11.columns[2].width = Inches(2.0)
tbl11.columns[3].width = Inches(1.6)
tbl11.columns[4].width = Inches(0.8)

h11 = ["ID", "Facility", "Required Items", "Status", "Age"]
for c_idx, h in enumerate(h11):
    cell = tbl11.cell(0, c_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = CARD_BG
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(9.5)
    p.font.color.rgb = CYAN

q_data = [
    ("REQ-0842", "Camp Alpha (Govt HSS)", "25 O2 Cylinders (CRITICAL)", "ACKNOWLEDGED", "14m"),
    ("REQ-0839", "Apollo Field Clinic", "40 Units O- Blood (CRITICAL)", "IN TRANSIT (15/40)", "42m"),
    ("REQ-0835", "Camp Delta (Maternity)", "500L Saline IV (HIGH)", "OFFER RECEIVED", "1h 10m"),
    ("REQ-0830", "PHC Kadambur", "1x 15kVA Generator (CRITICAL)", "VERIFIED RESOLVED", "2h 05m")
]

for r_idx, r in enumerate(q_data):
    for c_idx, val in enumerate(r):
        cell = tbl11.cell(r_idx + 1, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(14, 24, 43) if r_idx % 2 == 0 else CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.size = Pt(8.5)
        p.font.color.rgb = WHITE

# Right: Details Card
d_box = add_card(s11, Inches(8.5), Inches(3.7), Inches(4.033), Inches(2.9), CYAN)
tb_d = s11.shapes.add_textbox(Inches(8.65), Inches(3.8), Inches(3.733), Inches(2.7))
tf_d = tb_d.text_frame
tf_d.word_wrap = True
pt = tf_d.paragraphs[0]
pt.text = "DETAILS: REQ-0842"
pt.font.bold = True
pt.font.size = Pt(13)
pt.font.color.rgb = WHITE
pd = tf_d.add_paragraph()
pd.text = "Camp Alpha (Govt HSS) · Velachery Sector\nNeed: 25 D-Type Oxygen Cylinders (Pediatric Emergency)\n\nLINKED VOLUNTEER OFFER:\nCoimbatore Medical College Hospital\n15 Oxygen Cylinders available for immediate dispatch\n[APPROVE ALLOCATION]   [PARTIAL SPLIT: 15/25]\n\nAUDIT LOG: 14:02 Created · 14:08 Acknowledged · 14:15 Offer Logged"
pd.font.size = Pt(9.5)
pd.font.color.rgb = TEXT_MUTED
pd.space_before = Pt(4)

add_footer(s11, "SECTION D: COMMAND DASHBOARD")

# ==========================================
# SLIDE 12: DEMO SCENARIO
# ==========================================
s12 = prs.slides.add_slide(blank_layout)
set_slide_background(s12)
add_header(s12, "SECTION D · SIMULATED DRILL WALKTHROUGH", "Demo Scenario: Oxygen Supply Recovery in Cyclone Mandous", "Step-by-step resolution of a critical resource shortage amidst intermittent connectivity.", 12)

drill = [
    ("Phase 1 · 14:02\nOutage at Camp Beta", "Cyclone knocks out power & cellular links at Camp Beta. Dr. Priya drafts emergency request for 20 oxygen cylinders; saved in IndexedDB outbox.", RED),
    ("Phase 2 · 14:18\nTransient Signal Sync", "Emergency telecom van passes nearby, restoring 2G for 40s. EDITH sync manager replays with Idempotency Key idemp_91a8. Server commits REQ-0512.", CYAN),
    ("Phase 3 · 14:25\nICC Triage & Offer", "ICC Commander acknowledges REQ-0512. City Hospital logs voluntary offer: 12 cylinders available for immediate dispatch.", BLUE),
    ("Phase 4 · 14:34\nAtomic Lock & Transit", "ICC approves allocation. DB row-lock locks 12 units (0 over-allocation). Logistics vehicle TN-01-G-7712 departs with driver.", ORANGE),
    ("Phase 5 · 15:45\nVerified Receipt", "Driver drops cargo. Dr. Priya inspects and verifies 12 units. State splits: 12 confirmed closed, 8 units remain active in queue.", GREEN)
]

dr_w = Inches(2.2)
dr_h = Inches(3.6)
dr_top = Inches(1.8)

for idx, (title, desc, col) in enumerate(drill):
    dx = Inches(0.8 + idx * 2.4)
    add_card(s12, dx, dr_top, dr_w, dr_h, col)
    tb_dr = s12.shapes.add_textbox(dx + Inches(0.15), dr_top + Inches(0.15), dr_w - Inches(0.3), dr_h - Inches(0.3))
    tf_dr = tb_dr.text_frame
    tf_dr.word_wrap = True
    pt = tf_dr.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(11)
    pt.font.color.rgb = col
    pd = tf_dr.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = WHITE
    pd.space_before = Pt(6)

# Proof summary
sum_box = add_card(s12, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.1), CYAN)
tb_sum = s12.shapes.add_textbox(Inches(1.0), Inches(5.65), Inches(11.333), Inches(1.0))
tf_sum = tb_sum.text_frame
tf_sum.word_wrap = True
pt = tf_sum.paragraphs[0]
pt.text = "WHAT THIS DRILL PROVES"
pt.font.bold = True
pt.font.size = Pt(10)
pt.font.color.rgb = CYAN
pd = tf_sum.add_paragraph()
pd.text = "No data lost during blackout; zero duplicate tickets on transient reconnection; zero over-allocation occurred; and unfulfilled quota (8 cylinders) remained visible for secondary donor offers."
pd.font.size = Pt(10)
pd.font.color.rgb = WHITE
pd.space_before = Pt(3)

add_footer(s12, "SECTION D: SIMULATED DEMO DRILL")

# ==========================================
# SLIDE 13: CAPABILITIES & METRICS
# ==========================================
s13 = prs.slides.add_slide(blank_layout)
set_slide_background(s13)
add_header(s13, "SECTION D · VALUE PROPOSITION & METRICS", "Core Capabilities & Pilot Validation Metrics", "Translating architectural resilience into measurable life-saving coordination outcomes.", 13)

caps = [
    ("1. Operational Visibility", "• Unified multi-facility triage queue with urgency badges\n• Mandatory coordinator assignment eliminating orphaned tickets\n• Automated SLA timers tracking unacknowledged emergencies\n• Immutable audit history documenting every transition", CYAN),
    ("2. Resource Coordination", "• Sovereign voluntary offers without full inventory exposure\n• Transactional atomic allocation eliminating double-promising\n• Automatic state-splitting for partial fulfilments\n• Transit tracking and formal delivery dispute escalation", BLUE),
    ("3. Connectivity Resilience", "• Offline application shell caching via PWA Service Worker\n• IndexedDB transactional outbox preserving field drafts\n• Idempotent sync engine preventing duplicate tickets on retry\n• Two-phase recipient verification separating dispatch from delivery", GREEN)
]

cp_w = Inches(3.7)
cp_h = Inches(2.6)
cp_top = Inches(1.8)

for idx, (title, desc, col) in enumerate(caps):
    cpx = Inches(0.8 + idx * 4.0)
    add_card(s13, cpx, cp_top, cp_w, cp_h, col)
    tb_cp = s13.shapes.add_textbox(cpx + Inches(0.18), cp_top + Inches(0.18), cp_w - Inches(0.36), cp_h - Inches(0.36))
    tf_cp = tb_cp.text_frame
    tf_cp.word_wrap = True
    pt = tf_cp.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(13)
    pt.font.color.rgb = col
    pd = tf_cp.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = WHITE
    pd.space_before = Pt(6)

# Validation Hypotheses Box
h_box = add_card(s13, Inches(0.8), Inches(4.6), Inches(11.733), Inches(2.1), CYAN)
tb_h = s13.shapes.add_textbox(Inches(1.0), Inches(4.7), Inches(11.333), Inches(1.9))
tf_h = tb_h.text_frame
tf_h.word_wrap = True
pt = tf_h.paragraphs[0]
pt.text = "TARGET PILOT VALIDATION HYPOTHESES (MEASURABLE SUCCESS CRITERIA)"
pt.font.bold = True
pt.font.size = Pt(10.5)
pt.font.color.rgb = CYAN
pd = tf_h.add_paragraph()
pd.text = "• Time-to-Ack: >65% reduction in unacknowledged request duration vs informal channels\n• Named Accountability: 100% of emergency requests assigned to a designated human coordinator\n• Over-Allocation: 0 multi-allocation errors under high-concurrency disaster drills\n• Offline Replay: 100% successful replay of queued offline drafts upon reconnection\n• Verification: 100% of closed requests confirmed by recipient facility medical officer"
pd.font.size = Pt(10)
pd.font.color.rgb = WHITE
pd.space_before = Pt(4)

add_footer(s13, "SECTION D: CAPABILITIES & IMPACT")

# ==========================================
# SLIDE 14: BUSINESS MODEL & ADOPTION
# ==========================================
s14 = prs.slides.add_slide(blank_layout)
set_slide_background(s14)
add_header(s14, "SECTION E · SUSTAINABILITY & ADOPTION", "Institutional Adoption Model & Sustainability Pathway", "A viable deployment model for humanitarian agencies and public health authorities.", 14)

biz = [
    ("Target Ecosystem", "• State Disaster Management Authorities (SDMAs)\n• Public & Private Hospital Networks\n• Relief NGOs (Red Cross, MSF, Trusts)\n• District Emergency Operation Centres (DEOCs)", CYAN),
    ("Sustainability Model", "• Open-Core Public Good: Core platform open for trust\n• Institutional SLA & Hosting: Paid support & failover\n• Simulation Drills & Onboarding: Pre-disaster workshops\n• Zero Frontline Barrier: 100% free for temporary camps", BLUE),
    ("Adoption Pathway", "• Phase 1: Academic & Hackathon Pilot in Coimbatore\n• Phase 2: Tabletop simulation with local relief NGOs\n• Phase 3: Municipal flood mock-drill integration\n• Phase 4: Institutional deployment for active monsoons", GREEN)
]

bz_w = Inches(3.7)
bz_h = Inches(3.4)
bz_top = Inches(1.8)

for idx, (title, desc, col) in enumerate(biz):
    bx = Inches(0.8 + idx * 4.0)
    add_card(s14, bx, bz_top, bz_w, bz_h, col)
    tb_b = s14.shapes.add_textbox(bx + Inches(0.2), bz_top + Inches(0.2), bz_w - Inches(0.4), bz_h - Inches(0.4))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    pt = tf_b.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(13)
    pt.font.color.rgb = col
    pd = tf_b.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(10.5)
    pd.font.color.rgb = WHITE
    pd.space_before = Pt(8)

# Mitigation
f_box = add_card(s14, Inches(0.8), Inches(5.4), Inches(11.733), Inches(1.3), CYAN)
tb_f = s14.shapes.add_textbox(Inches(1.0), Inches(5.45), Inches(11.333), Inches(1.2))
tf_f = tb_f.text_frame
tf_f.word_wrap = True
pt = tf_f.paragraphs[0]
pt.text = "MITIGATING PUBLIC SECTOR ADOPTION FRICTION"
pt.font.bold = True
pt.font.size = Pt(10)
pt.font.color.rgb = CYAN
pd = tf_f.add_paragraph()
pd.text = "Avoiding long procurement delays by providing a turnkey Docker image and zero-login guest drafting for emergency field responders during acute crises."
pd.font.size = Pt(10)
pd.font.color.rgb = WHITE
pd.space_before = Pt(4)

add_footer(s14, "SECTION E: BUSINESS MODEL & ADOPTION")

# ==========================================
# SLIDE 15: FEASIBILITY & RISK MITIGATION
# ==========================================
s15 = prs.slides.add_slide(blank_layout)
set_slide_background(s15)
add_header(s15, "SECTION E · FEASIBILITY & RISK MITIGATION", "Technical Feasibility, Testing Matrix & Risk Safeguards", "Rigorous software engineering that separates practical innovation from speculative concepts.", 15)

feas_w = Inches(5.7)
feas_h = Inches(3.3)
feas_top = Inches(1.8)

# Left: Testing
add_card(s15, Inches(0.8), feas_top, feas_w, feas_h, CYAN)
tb_ts = s15.shapes.add_textbox(Inches(1.0), feas_top + Inches(0.2), feas_w - Inches(0.4), feas_h - Inches(0.4))
tf_ts = tb_ts.text_frame
tf_ts.word_wrap = True
pt = tf_ts.paragraphs[0]
pt.text = "COMPREHENSIVE TESTING MATRIX"
pt.font.bold = True
pt.font.size = Pt(13)
pt.font.color.rgb = CYAN
pd = tf_ts.add_paragraph()
pd.text = "• Synthetic Disaster Load: Scripted simulation of 100+ concurrent requests across 20 camps\n• Concurrency Race Tests: Pytest suite firing 10 simultaneous allocations against 1 stock pool\n• Network Chaos Verification: Simulated connection drops during active outbox sync\n• State Transition Coverage: 100% automated branch testing across all 7 lifecycle states"
pd.font.size = Pt(10)
pd.font.color.rgb = WHITE
pd.space_before = Pt(8)

# Right: Risk vs Mitigation
add_card(s15, Inches(6.8), feas_top, feas_w, feas_h, ORANGE)
tb_rk = s15.shapes.add_textbox(Inches(7.0), feas_top + Inches(0.2), feas_w - Inches(0.4), feas_h - Inches(0.4))
tf_rk = tb_rk.text_frame
tf_rk.word_wrap = True
pt = tf_rk.paragraphs[0]
pt.text = "OPERATIONAL RISK VS. MITIGATION"
pt.font.bold = True
pt.font.size = Pt(13)
pt.font.color.rgb = ORANGE
pd = tf_rk.add_paragraph()
pd.text = "• Risk: Stale Resource Information\n  Mitigation: Request-centric model; facilities volunteer only verified surplus.\n• Risk: Multi-Device Conflict & Overwrites\n  Mitigation: Optimistic locking with version integers; server rejects stale updates.\n• Risk: High-Stress Field Adoption Resistance\n  Mitigation: Ultra-lightweight form design requiring under 30s to submit."
pd.font.size = Pt(10)
pd.font.color.rgb = WHITE
pd.space_before = Pt(8)

# Bottom hardware box
hw_box = add_card(s15, Inches(0.8), Inches(5.3), Inches(11.733), Inches(1.4), GREEN)
tb_hw = s15.shapes.add_textbox(Inches(1.0), Inches(5.35), Inches(11.333), Inches(1.3))
tf_hw = tb_hw.text_frame
tf_hw.word_wrap = True
pt = tf_hw.paragraphs[0]
pt.text = "FIELD DEPLOYMENT FEASIBILITY"
pt.font.bold = True
pt.font.size = Pt(10)
pt.font.color.rgb = GREEN
pd = tf_hw.add_paragraph()
pd.text = "• Hardware: Runs smoothly on a $20/mo cloud VM or a local command laptop\n• Zero Licensing: 100% open standards (FastAPI, PostgreSQL, Vanilla PWA)\n• Universal Support: Operates on Android, iOS, Windows, Linux, and macOS without app store installs"
pd.font.size = Pt(10)
pd.font.color.rgb = WHITE
pd.space_before = Pt(4)

add_footer(s15, "SECTION E: FEASIBILITY & RISK MITIGATION")

# ==========================================
# SLIDE 16: ROADMAP & CLOSING
# ==========================================
s16 = prs.slides.add_slide(blank_layout)
set_slide_background(s16)
add_header(s16, "SECTION F · ROADMAP & CLOSING VISION", "Execution Roadmap & The EDITH Commitment", "From hackathon baseline to hardened institutional disaster deployment.", 16)

phases = [
    ("Phase 1\nCore Coordination", "Auth, RBAC, request lifecycle state machine, & audit logs.\n[BASELINE READY]", GREEN),
    ("Phase 2\nResource Flow", "Voluntary offers, atomic row locks, & verified fulfilment.\n[BASELINE READY]", GREEN),
    ("Phase 3\nOffline Resilience", "Service Worker cache, IndexedDB outbox, & sync engine.\n[BASELINE READY]", CYAN),
    ("Phase 4\nHardening & QA", "Chaos testing, concurrency race stress tests, & mock drills.\n[PLANNED Q1]", BLUE),
    ("Phase 5\nAssistive AI", "Speech-to-text intake, request summaries, & triage advice.\n[FUTURE SCOPE]", ORANGE)
]

ph_w = Inches(2.2)
ph_h = Inches(2.5)
ph_top = Inches(1.8)

for idx, (title, desc, col) in enumerate(phases):
    px = Inches(0.8 + idx * 2.4)
    add_card(s16, px, ph_top, ph_w, ph_h, col)
    tb_ph = s16.shapes.add_textbox(px + Inches(0.12), ph_top + Inches(0.12), ph_w - Inches(0.24), ph_h - Inches(0.24))
    tf_ph = tb_ph.text_frame
    tf_ph.word_wrap = True
    pt = tf_ph.paragraphs[0]
    pt.text = title
    pt.font.bold = True
    pt.font.size = Pt(11)
    pt.font.color.rgb = col
    pd = tf_ph.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(9)
    pd.font.color.rgb = WHITE
    pd.space_before = Pt(6)

# Closing Statement Card
cl_box = add_card(s16, Inches(0.8), Inches(4.55), Inches(11.733), Inches(2.15), CYAN)
tb_cl = s16.shapes.add_textbox(Inches(1.0), Inches(4.65), Inches(11.333), Inches(1.95))
tf_cl = tb_cl.text_frame
tf_cl.word_wrap = True

pt = tf_cl.paragraphs[0]
pt.text = '"EDITH turns fragmented emergency requests into accountable, traceable coordination — designed for the moments connectivity cannot be taken for granted."'
pt.font.bold = True
pt.font.italic = True
pt.font.size = Pt(16)
pt.font.color.rgb = WHITE
pt.alignment = PP_ALIGN.CENTER

p_team = tf_cl.add_paragraph()
p_team.text = "TEAM JARVIS · RATHINAM TECHNICAL CAMPUS"
p_team.font.bold = True
p_team.font.size = Pt(12)
p_team.font.color.rgb = CYAN
p_team.alignment = PP_ALIGN.CENTER
p_team.space_before = Pt(12)

p_names = tf_cl.add_paragraph()
p_names.text = "Purusouthanan B (Lead) · Sakthivel R · Sanjo Thomas · Sheik Arsh S"
p_names.font.size = Pt(11)
p_names.font.color.rgb = TEXT_MUTED
p_names.alignment = PP_ALIGN.CENTER
p_names.space_before = Pt(4)

add_footer(s16, "SECTION F: ROADMAP & CLOSING")

prs.save(PPTX_PATH)
print(f"Success! PowerPoint presentation saved to: {PPTX_PATH} ({round(os.path.getsize(PPTX_PATH)/1024, 1)} KB)")
